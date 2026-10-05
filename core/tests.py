import json

from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
from django.utils import timezone

from core.models import (
    TestPlan, TestStep, TestRun, RunStepResult, Incident, Finding,
)
from rest_framework_api_key.models import APIKey


class TestPlanModelTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')

    def test_create_test_plan(self):
        plan = TestPlan.objects.create(
            name='Login Flow',
            project_name='WebApp',
            test_scope='User authentication',
            created_by=self.user,
        )
        self.assertEqual(plan.name, 'Login Flow')
        self.assertEqual(plan.project_name, 'WebApp')
        self.assertEqual(plan.created_by, self.user)
        self.assertEqual(plan.total_steps, 0)
        self.assertIsNone(plan.latest_run)
        self.assertEqual(str(plan), 'Login Flow')

    def test_test_plan_timestamps(self):
        plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.assertIsNotNone(plan.created_at)
        self.assertIsNotNone(plan.updated_at)

    def test_plan_type_defaults_to_qa(self):
        plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.assertEqual(plan.plan_type, 'qa')

    def test_plan_type_security(self):
        plan = TestPlan.objects.create(
            name='Pentest',
            plan_type='security',
            created_by=self.user,
        )
        self.assertEqual(plan.plan_type, 'security')
        self.assertEqual(str(plan), 'Pentest')


class TestStepModelTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)

    def test_create_test_step(self):
        step = TestStep.objects.create(
            plan=self.plan,
            name='Navigate to login',
            action_description='Open the login page',
            expected_outcome='Login form is displayed',
            order_index=0,
        )
        self.assertEqual(step.name, 'Navigate to login')
        self.assertEqual(step.order_index, 0)
        self.assertTrue(step.active)
        self.assertEqual(str(step), '[Plan] Navigate to login')

    def test_step_ordering(self):
        TestStep.objects.create(plan=self.plan, name='Step 2', order_index=2, action_description='x', expected_outcome='x')
        TestStep.objects.create(plan=self.plan, name='Step 1', order_index=1, action_description='x', expected_outcome='x')
        TestStep.objects.create(plan=self.plan, name='Step 0', order_index=0, action_description='x', expected_outcome='x')
        steps = list(TestStep.objects.filter(plan=self.plan))
        self.assertEqual(steps[0].name, 'Step 0')
        self.assertEqual(steps[1].name, 'Step 1')
        self.assertEqual(steps[2].name, 'Step 2')

    def test_total_steps_property(self):
        TestStep.objects.create(plan=self.plan, name='A', order_index=0, active=True, action_description='x', expected_outcome='x')
        TestStep.objects.create(plan=self.plan, name='B', order_index=1, active=False, action_description='x', expected_outcome='x')
        self.assertEqual(self.plan.total_steps, 1)


class TestRunModelTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.step1 = TestStep.objects.create(plan=self.plan, name='Step 1', order_index=0, active=True, action_description='x', expected_outcome='x')
        self.step2 = TestStep.objects.create(plan=self.plan, name='Step 2', order_index=1, active=True, action_description='x', expected_outcome='x')

    def test_create_test_run(self):
        run = TestRun.objects.create(plan=self.plan, agent_id='claude-code')
        self.assertEqual(run.status, 'pending')
        self.assertIsNone(run.completed_at)
        self.assertEqual(run.agent_id, 'claude-code')
        self.assertEqual(str(run), f'Run #{run.id} - Plan (Pending)')

    def test_run_step_counts(self):
        run = TestRun.objects.create(plan=self.plan)
        self.assertEqual(run.total_steps, 2)
        self.assertEqual(run.passed_steps, 0)
        self.assertEqual(run.failed_steps, 0)
        self.assertEqual(run.skipped_steps, 0)
        self.assertEqual(run.pending_steps, 2)

        RunStepResult.objects.create(run=run, step=self.step1, status='passed')
        RunStepResult.objects.create(run=run, step=self.step2, status='failed')
        run.refresh_from_db()
        self.assertEqual(run.passed_steps, 1)
        self.assertEqual(run.failed_steps, 1)
        self.assertEqual(run.pending_steps, 0)

    def test_complete_run(self):
        run = TestRun.objects.create(plan=self.plan, status='running')
        run.status = 'completed'
        run.completed_at = timezone.now()
        run.save()
        self.assertEqual(run.status, 'completed')
        self.assertIsNotNone(run.completed_at)


class RunStepResultModelTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.step = TestStep.objects.create(plan=self.plan, name='Step', order_index=0, action_description='x', expected_outcome='x')
        self.run = TestRun.objects.create(plan=self.plan)

    def test_create_result(self):
        result = RunStepResult.objects.create(
            run=self.run,
            step=self.step,
            status='passed',
            log_message='Everything worked',
        )
        self.assertEqual(result.status, 'passed')
        self.assertEqual(result.log_message, 'Everything worked')
        self.assertEqual(str(result), 'Step → Passed')

    def test_unique_constraint(self):
        RunStepResult.objects.create(run=self.run, step=self.step, status='passed')
        with self.assertRaises(Exception):
            RunStepResult.objects.create(run=self.run, step=self.step, status='failed')


class IncidentModelTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.step = TestStep.objects.create(plan=self.plan, name='Step', order_index=0, action_description='x', expected_outcome='x')
        self.run = TestRun.objects.create(plan=self.plan)
        self.result = RunStepResult.objects.create(run=self.run, step=self.step, status='failed')

    def test_create_incident(self):
        incident = Incident.objects.create(
            run_step_result=self.result,
            summary='Button not found',
            reproduction_steps='1. Go to page\n2. Click button',
            severity='high',
        )
        self.assertEqual(incident.summary, 'Button not found')
        self.assertEqual(incident.severity, 'high')
        self.assertFalse(incident.resolved)
        self.assertIsNone(incident.assigned_to)
        self.assertEqual(str(incident), 'Incident #1: Button not found')

    def test_severity_choices(self):
        for severity in ('low', 'medium', 'high'):
            incident = Incident.objects.create(
                run_step_result=self.result,
                summary=f'{severity} issue',
                reproduction_steps='steps',
                severity=severity,
            )
            self.assertEqual(incident.severity, severity)

    def test_resolve_incident(self):
        incident = Incident.objects.create(
            run_step_result=self.result,
            summary='Bug',
            reproduction_steps='steps',
        )
        self.assertFalse(incident.resolved)
        incident.resolved = True
        incident.save()
        self.assertTrue(incident.resolved)


class TestPlanSerializerTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client = Client()
        self.client.login(username='testuser', password='testpass')

    def test_create_test_plan(self):
        response = self.client.post(
            reverse('api:testplan-list'),
            {'name': 'New Plan', 'project_name': 'Project', 'test_scope': 'Scope'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(TestPlan.objects.count(), 1)

    def test_list_test_plans(self):
        TestPlan.objects.create(name='Plan 1', created_by=self.user)
        TestPlan.objects.create(name='Plan 2', created_by=self.user)
        response = self.client.get(reverse('api:testplan-list'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 2)

    def test_create_security_plan(self):
        response = self.client.post(
            reverse('api:testplan-list'),
            {'name': 'Pentest', 'plan_type': 'security', 'test_scope': 'http://localhost:8000'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        plan = TestPlan.objects.get(pk=response.json()['id'])
        self.assertEqual(plan.plan_type, 'security')

    def test_filter_plans_by_plan_type(self):
        TestPlan.objects.create(name='QA Plan', plan_type='qa', created_by=self.user)
        TestPlan.objects.create(name='Security Plan', plan_type='security', created_by=self.user)
        response = self.client.get(reverse('api:testplan-list'), {'plan_type': 'security'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 1)
        self.assertEqual(response.json()['results'][0]['name'], 'Security Plan')

    def test_filter_plans_by_qa_type(self):
        TestPlan.objects.create(name='QA Plan', plan_type='qa', created_by=self.user)
        TestPlan.objects.create(name='Security Plan', plan_type='security', created_by=self.user)
        response = self.client.get(reverse('api:testplan-list'), {'plan_type': 'qa'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 1)
        self.assertEqual(response.json()['results'][0]['name'], 'QA Plan')


class TestStepSerializerTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.client = Client()
        self.client.login(username='testuser', password='testpass')

    def test_create_test_step(self):
        response = self.client.post(
            reverse('api:teststep-list'),
            {
                'plan': self.plan.id,
                'name': 'Step 1',
                'action_description': 'Do something',
                'expected_outcome': 'Something happens',
                'order_index': 0,
            },
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(TestStep.objects.count(), 1)

    def test_filter_steps_by_plan(self):
        TestStep.objects.create(plan=self.plan, name='A', order_index=0, action_description='x', expected_outcome='x')
        other_plan = TestPlan.objects.create(name='Other', created_by=self.user)
        TestStep.objects.create(plan=other_plan, name='B', order_index=0, action_description='x', expected_outcome='x')
        response = self.client.get(reverse('api:teststep-list'), {'plan': self.plan.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 1)


class TestRunSerializerTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.client = Client()
        self.client.login(username='testuser', password='testpass')

    def test_create_test_run(self):
        response = self.client.post(
            reverse('api:testrun-list'),
            {'plan': self.plan.id, 'agent_id': 'test-agent'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        run = TestRun.objects.first()
        self.assertEqual(run.status, 'pending')
        self.assertEqual(run.agent_id, 'test-agent')

    def test_complete_test_run(self):
        run = TestRun.objects.create(plan=self.plan, status='running')
        response = self.client.post(
            reverse('api:testrun-complete', kwargs={'pk': run.id}),
            {'status': 'completed'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        run.refresh_from_db()
        self.assertEqual(run.status, 'completed')
        self.assertIsNotNone(run.completed_at)


class AuthTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client = Client()

    def test_login_required(self):
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response.url)

    def test_login_success(self):
        self.client.login(username='testuser', password='testpass')
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)

    def test_unauthenticated_api_access_denied(self):
        response = self.client.get(reverse('api:testplan-list'))
        self.assertEqual(response.status_code, 403)


class FirstTimeLaunchTest(TestCase):
    """Tests for the first-time launch / registration flow."""

    def test_root_redirects_to_register_when_no_users(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('register'))

    def test_login_redirects_to_register_when_no_users(self):
        response = self.client.get(reverse('login'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('register'))

    def test_login_shows_form_when_users_exist(self):
        User.objects.create_user(username='testuser', password='testpass')
        response = self.client.get(reverse('login'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'login.html')

    def test_register_page_loads_when_no_users(self):
        response = self.client.get(reverse('register'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'register.html')

    def test_register_redirects_to_login_when_users_exist(self):
        User.objects.create_user(username='testuser', password='testpass')
        response = self.client.get(reverse('register'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('login'))

    def test_register_creates_user_and_logs_in(self):
        response = self.client.post(reverse('register'), {
            'username': 'admin',
            'password': 'strongpassword',
            'password_confirm': 'strongpassword',
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(User.objects.count(), 1)
        user = User.objects.first()
        self.assertEqual(user.username, 'admin')
        self.assertTrue(user.is_superuser)
        self.assertTrue(user.is_staff)

    def test_register_requires_password_match(self):
        response = self.client.post(reverse('register'), {
            'username': 'admin',
            'password': 'password1',
            'password_confirm': 'password2',
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(User.objects.count(), 0)
        self.assertIn('Passwords do not match', str(response.context['errors']))

    def test_register_requires_username(self):
        response = self.client.post(reverse('register'), {
            'username': '',
            'password': 'password1',
            'password_confirm': 'password1',
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(User.objects.count(), 0)
        self.assertIn('Username is required.', str(response.context['errors']))

    def test_register_requires_password(self):
        response = self.client.post(reverse('register'), {
            'username': 'admin',
            'password': '',
            'password_confirm': '',
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(User.objects.count(), 0)

    def test_register_redirects_to_dashboard_after_success(self):
        response = self.client.post(reverse('register'), {
            'username': 'admin',
            'password': 'strongpassword',
            'password_confirm': 'strongpassword',
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('dashboard'))

    def test_register_second_user_blocked(self):
        User.objects.create_user(username='first', password='password')
        response = self.client.post(reverse('register'), {
            'username': 'second',
            'password': 'strongpassword',
            'password_confirm': 'strongpassword',
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('login'))
        self.assertEqual(User.objects.count(), 1)


class APIKeyManagementTest(TestCase):
    """Tests for the API key management endpoints."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username='admin', password='adminpass', is_staff=True, is_superuser=True,
        )
        self.regular_user = User.objects.create_user(
            username='regular', password='regularpass',
        )
        self.client = Client()

    # --- List ---

    def test_list_api_keys_requires_auth(self):
        response = self.client.get(reverse('api:apikey-list'))
        self.assertEqual(response.status_code, 403)

    def test_list_api_keys_admin(self):
        self.client.login(username='admin', password='adminpass')
        APIKey.objects.create_key(name='test-key')
        response = self.client.get(reverse('api:apikey-list'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 1)

    def test_list_api_keys_regular_user_denied(self):
        self.client.login(username='regular', password='regularpass')
        response = self.client.get(reverse('api:apikey-list'))
        self.assertEqual(response.status_code, 403)

    def test_list_api_keys_filter_revoked(self):
        self.client.login(username='admin', password='adminpass')
        APIKey.objects.create_key(name='active-key')
        instance, _ = APIKey.objects.create_key(name='revoked-key')
        instance.revoked = True
        instance.save()
        response = self.client.get(reverse('api:apikey-list'), {'revoked': 'false'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 1)
        self.assertEqual(response.json()['results'][0]['name'], 'active-key')

    # --- Create ---

    def test_create_api_key_requires_auth(self):
        response = self.client.post(
            reverse('api:apikey-list'),
            {'name': 'new-key'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 403)

    def test_create_api_key_regular_user_denied(self):
        self.client.login(username='regular', password='regularpass')
        response = self.client.post(
            reverse('api:apikey-list'),
            {'name': 'new-key'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 403)

    def test_create_api_key_success(self):
        self.client.login(username='admin', password='adminpass')
        response = self.client.post(
            reverse('api:apikey-list'),
            {'name': 'claude-code-agent'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['name'], 'claude-code-agent')
        self.assertIn('key', data)
        self.assertIn('prefix', data)
        self.assertFalse(data['revoked'])
        # Raw key should be a non-empty string
        self.assertIsInstance(data['key'], str)
        self.assertTrue(len(data['key']) > 10)

    def test_create_api_key_requires_name(self):
        self.client.login(username='admin', password='adminpass')
        response = self.client.post(
            reverse('api:apikey-list'),
            {},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_create_api_key_with_expiry(self):
        self.client.login(username='admin', password='adminpass')
        response = self.client.post(
            reverse('api:apikey-list'),
            {'name': 'temp-key', 'expiry_date': '2026-12-31'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['expiry_date'][:10], '2026-12-31')

    # --- Revoke ---

    def test_revoke_api_key(self):
        self.client.login(username='admin', password='adminpass')
        instance, raw_key = APIKey.objects.create_key(name='to-revoke')
        response = self.client.post(
            reverse('api:apikey-revoke', kwargs={'prefix': instance.prefix}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['revoked'])
        instance.refresh_from_db()
        self.assertTrue(instance.revoked)

    def test_revoke_api_key_not_found(self):
        self.client.login(username='admin', password='adminpass')
        response = self.client.post(
            reverse('api:apikey-revoke', kwargs={'prefix': 'nonexistent999'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 404)

    # --- Raw key not exposed on list ---

    def test_raw_key_not_in_list_response(self):
        self.client.login(username='admin', password='adminpass')
        APIKey.objects.create_key(name='secret-key')
        response = self.client.get(reverse('api:apikey-list'))
        self.assertEqual(response.status_code, 200)
        for item in response.json()['results']:
            self.assertNotIn('key', item)

    # --- UI view ---

    def test_api_keys_ui_requires_login(self):
        response = self.client.get('/api-keys/')
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response.url)

    def test_api_keys_ui_success(self):
        self.client.login(username='admin', password='adminpass')
        response = self.client.get('/api-keys/')
        self.assertEqual(response.status_code, 200)


class BulkCreateTestStepsTest(TestCase):
    """Tests for the bulk_create_test_steps endpoint."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.client = Client()
        self.client.login(username='testuser', password='testpass')

    def test_bulk_create_steps_success(self):
        payload = {
            'plan': self.plan.id,
            'steps': [
                {
                    'name': 'Step A',
                    'action_description': 'Do A',
                    'expected_outcome': 'A happens',
                },
                {
                    'name': 'Step B',
                    'action_description': 'Do B',
                    'expected_outcome': 'B happens',
                    'preconditions': 'Step A passed',
                },
                {
                    'name': 'Step C',
                    'action_description': 'Do C',
                    'expected_outcome': 'C happens',
                },
            ],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['created'], 3)
        self.assertEqual(len(data['steps']), 3)

        # Verify steps were created with correct order
        steps = TestStep.objects.filter(plan=self.plan).order_by('order_index')
        self.assertEqual(steps.count(), 3)
        self.assertEqual(steps[0].name, 'Step A')
        self.assertEqual(steps[1].name, 'Step B')
        self.assertEqual(steps[2].name, 'Step C')

    def test_bulk_create_auto_order_index(self):
        # Create one step first
        TestStep.objects.create(
            plan=self.plan, name='Existing', order_index=5,
            action_description='x', expected_outcome='x',
        )
        payload = {
            'plan': self.plan.id,
            'steps': [
                {'name': 'A', 'action_description': 'a', 'expected_outcome': 'a'},
                {'name': 'B', 'action_description': 'b', 'expected_outcome': 'b'},
            ],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        data = response.json()
        # Should start at order_index 6 (after existing max of 5)
        self.assertEqual(data['steps'][0]['order_index'], 6)
        self.assertEqual(data['steps'][1]['order_index'], 7)

    def test_bulk_create_with_explicit_order_index(self):
        payload = {
            'plan': self.plan.id,
            'steps': [
                {'name': 'A', 'action_description': 'a', 'expected_outcome': 'a', 'order_index': 10},
                {'name': 'B', 'action_description': 'b', 'expected_outcome': 'b', 'order_index': 20},
            ],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        steps = TestStep.objects.filter(plan=self.plan).order_by('order_index')
        self.assertEqual(steps[0].order_index, 10)
        self.assertEqual(steps[1].order_index, 20)

    def test_bulk_create_missing_required_field(self):
        payload = {
            'plan': self.plan.id,
            'steps': [
                {'name': 'Incomplete'},  # missing action_description and expected_outcome
            ],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_create_invalid_plan(self):
        payload = {
            'plan': 99999,
            'steps': [
                {'name': 'A', 'action_description': 'a', 'expected_outcome': 'a'},
            ],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_create_empty_steps(self):
        payload = {
            'plan': self.plan.id,
            'steps': [],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_create_includes_all_fields(self):
        payload = {
            'plan': self.plan.id,
            'steps': [
                {
                    'name': 'Full',
                    'action_description': 'Full action',
                    'expected_outcome': 'Full outcome',
                    'preconditions': 'Precondition met',
                    'active': False,
                },
            ],
        }
        response = self.client.post(
            reverse('api:teststep-bulk-create'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        step_data = response.json()['steps'][0]
        self.assertEqual(step_data['name'], 'Full')
        self.assertEqual(step_data['preconditions'], 'Precondition met')
        self.assertFalse(step_data['active'])
        self.assertIn('id', step_data)
        self.assertIn('created_at', step_data)


class BulkLogStepResultsTest(TestCase):
    """Tests for the bulk_log_step_results endpoint."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.step1 = TestStep.objects.create(
            plan=self.plan, name='Step 1', order_index=0,
            action_description='x', expected_outcome='x',
        )
        self.step2 = TestStep.objects.create(
            plan=self.plan, name='Step 2', order_index=1,
            action_description='x', expected_outcome='x',
        )
        self.step3 = TestStep.objects.create(
            plan=self.plan, name='Step 3', order_index=2,
            action_description='x', expected_outcome='x',
        )
        self.run = TestRun.objects.create(plan=self.plan)
        self.client = Client()
        self.client.login(username='testuser', password='testpass')

    def test_bulk_log_results_success(self):
        payload = {
            'run': self.run.id,
            'results': [
                {'step': self.step1.id, 'status': 'passed', 'log_message': 'OK'},
                {'step': self.step2.id, 'status': 'failed', 'log_message': 'Error'},
                {'step': self.step3.id, 'status': 'skipped'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['created'], 3)
        self.assertEqual(len(data['skipped']), 0)
        self.assertEqual(len(data['results']), 3)

        # Verify in database
        self.assertEqual(RunStepResult.objects.filter(run=self.run).count(), 3)

    def test_bulk_log_skips_existing(self):
        # Create one result first
        RunStepResult.objects.create(
            run=self.run, step=self.step1, status='passed',
        )
        payload = {
            'run': self.run.id,
            'results': [
                {'step': self.step1.id, 'status': 'failed', 'log_message': 'Should be skipped'},
                {'step': self.step2.id, 'status': 'passed'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['created'], 1)
        self.assertEqual(len(data['skipped']), 1)
        self.assertEqual(data['skipped'][0]['reason'], 'Result already exists')

        # Original result should be unchanged
        result = RunStepResult.objects.get(run=self.run, step=self.step1)
        self.assertEqual(result.status, 'passed')

    def test_bulk_log_skips_missing_step(self):
        payload = {
            'run': self.run.id,
            'results': [
                {'step': 99999, 'status': 'passed'},  # non-existent step
                {'step': self.step1.id, 'status': 'passed'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['created'], 1)
        self.assertEqual(len(data['skipped']), 1)
        self.assertEqual(data['skipped'][0]['reason'], 'Step not found')

    def test_bulk_log_invalid_status(self):
        payload = {
            'run': self.run.id,
            'results': [
                {'step': self.step1.id, 'status': 'invalid_status'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_log_missing_required_field(self):
        payload = {
            'run': self.run.id,
            'results': [
                {'step': self.step1.id},  # missing status
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_log_invalid_run(self):
        payload = {
            'run': 99999,
            'results': [
                {'step': self.step1.id, 'status': 'passed'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_log_empty_results(self):
        payload = {
            'run': self.run.id,
            'results': [],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_log_includes_result_fields(self):
        payload = {
            'run': self.run.id,
            'results': [
                {'step': self.step1.id, 'status': 'passed', 'log_message': 'Detail log'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        result_data = response.json()['results'][0]
        self.assertEqual(result_data['status'], 'passed')
        self.assertEqual(result_data['log_message'], 'Detail log')
        self.assertEqual(result_data['step_name'], 'Step 1')
        self.assertIn('id', result_data)
        self.assertIn('created_at', result_data)

    def test_bulk_log_all_statuses(self):
        payload = {
            'run': self.run.id,
            'results': [
                {'step': self.step1.id, 'status': 'passed'},
                {'step': self.step2.id, 'status': 'failed'},
                {'step': self.step3.id, 'status': 'skipped'},
            ],
        }
        response = self.client.post(
            reverse('api:runstepresult-bulk-log'),
            payload,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['created'], 3)


class SkipStepsTest(TestCase):
    """Tests for POST /api/test-runs/{id}/skip-steps/."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        for i in range(12):
            section = 'Alpha' if i < 6 else 'Beta'
            TestStep.objects.create(
                plan=self.plan, name=f'Step {i + 1}', order_index=i,
                action_description='x', expected_outcome='x', section=section,
            )
        self.run = TestRun.objects.create(plan=self.plan, agent_id='skip-test')
        self.client = Client()
        self.client.login(username='testuser', password='testpass')

    def _skip(self, payload):
        return self.client.post(
            reverse('api:testrun-skip-steps', args=[self.run.id]),
            payload, content_type='application/json',
        )

    def test_skip_range_creates_skipped_results(self):
        r = self._skip({'ranges': [[1, 6]]})
        self.assertEqual(r.status_code, 201)
        d = r.json()
        self.assertEqual(d['skipped'], 6)
        self.assertEqual(d['unchanged'], 0)
        self.assertEqual(d['not_found_positions'], [])
        self.assertEqual(RunStepResult.objects.filter(run=self.run, status='skipped').count(), 6)

    def test_skip_is_idempotent(self):
        self._skip({'ranges': [[1, 6]]})
        d = self._skip({'ranges': [[1, 6]]}).json()
        self.assertEqual(d['skipped'], 0)
        self.assertEqual(d['unchanged'], 6)
        self.assertEqual(RunStepResult.objects.filter(run=self.run, step__section='Alpha').count(), 6)

    def test_skip_preserves_existing_results(self):
        step7 = TestStep.objects.filter(plan=self.plan, order_index=6).first()
        RunStepResult.objects.create(run=self.run, step=step7, status='passed')
        d = self._skip({'ranges': [[7, 8]]}).json()
        self.assertEqual(d['skipped'], 1)
        self.assertEqual(d['unchanged'], 1)

    def test_skip_out_of_range_reported_not_errored(self):
        d = self._skip({'ranges': [[1, 15]]}).json()
        self.assertEqual(d['not_found_positions'], [13, 14, 15])
        self.assertEqual(d['skipped'], 12)

    def test_skip_multiple_ranges(self):
        d = self._skip({'ranges': [[1, 2], [11, 12]]}).json()
        self.assertEqual(d['skipped'], 4)

    def test_skip_exclude_section(self):
        d = self._skip({'exclude_section': 'Alpha'}).json()
        self.assertEqual(d['skipped'], 6)  # all Beta steps
        self.assertFalse(RunStepResult.objects.filter(run=self.run, step__section='Alpha').exists())

    def test_skip_validation_errors(self):
        for payload in ({'ranges': [[5, 2]]}, {'ranges': []}, {}, {'ranges': [['a', 2]]}):
            self.assertEqual(self._skip(payload).status_code, 400)

    def test_skip_requires_auth(self):
        self.client.logout()
        r = self.client.post(reverse('api:testrun-skip-steps', args=[self.run.id]),
                             {'ranges': [[1, 1]]}, content_type='application/json')
        self.assertEqual(r.status_code, 403)


class McpSectionTest(TestCase):
    """The MCP layer must forward 'section' on step creation."""

    def _call_create(self, **kwargs):
        import mcp_server.server as m

        captured = {}

        class FakeResp:
            def raise_for_status(self):
                pass

            def json(self):
                return {'id': 1, 'section': 'Auth'}

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def post(self, url, json=None):
                captured['url'] = url
                captured['json'] = json
                return FakeResp()

        original = m._client
        m._client = lambda: FakeClient()
        try:
            m.create_test_step(**kwargs)
        finally:
            m._client = original

        return captured

    def test_create_test_step_forwards_section(self):
        captured = self._call_create(
            plan_id=1, name='S', action_description='a',
            expected_outcome='b', section='Auth',
        )
        self.assertEqual(captured['url'], '/api/test-steps/')
        self.assertEqual(captured['json']['section'], 'Auth')

    def test_create_test_step_without_section_omits_it(self):
        captured = self._call_create(
            plan_id=1, name='S', action_description='a', expected_outcome='b',
        )
        self.assertNotIn('section', captured['json'])

    def test_rest_single_create_persists_section(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='Plan', created_by=self.user)
        self.client.login(username='testuser', password='testpass')
        r = self.client.post(
            reverse('api:teststep-list'),
            {'plan': self.plan.id, 'name': 'S', 'action_description': 'a',
             'expected_outcome': 'b', 'section': 'Auth'},
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.json()['section'], 'Auth')




class McpDocstringsTest(TestCase):
    """MCP tool docstrings must document valid enum values and JSON examples."""

    def test_create_finding_documents_categories(self):
        import mcp_server.server as m
        doc = m.create_finding.__doc__
        for cat in ("info", "suggestion", "recommendation", "critical"):
            self.assertIn(cat, doc)

    def test_log_step_result_documents_statuses(self):
        import mcp_server.server as m
        doc = m.log_step_result.__doc__
        for st in ("passed", "failed", "skipped"):
            self.assertIn(st, doc)

    def test_bulk_log_documents_json_example(self):
        import mcp_server.server as m
        doc = m.bulk_log_step_results.__doc__
        self.assertIn('"step"', doc)
        self.assertIn('"status"', doc)

    def test_bulk_create_documents_json_example_with_section(self):
        import mcp_server.server as m
        doc = m.bulk_create_test_steps.__doc__
        for key in ('"name"', '"action_description"', '"expected_outcome"', '"section"'):
            self.assertIn(key, doc)


class McpPaginationTest(TestCase):
    """get_step_results must return ALL results by default (page-walk)."""

    PAGE1 = {
        "count": 4,
        "next": "http://t/api/step-results/?page=2",
        "previous": None,
        "results": [{"id": 1, "step": 1}, {"id": 2, "step": 2}],
    }
    PAGE2 = {
        "count": 4,
        "next": None,
        "previous": "http://t/api/step-results/?page=1",
        "results": [{"id": 3, "step": 3}, {"id": 4, "step": 4}],
    }

    def _call(self, pages, **kwargs):
        import mcp_server.server as m

        calls = []

        class FakeResp:
            def __init__(self, data):
                self._data = data

            def raise_for_status(self):
                pass

            def json(self):
                return self._data

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get(self, url, params=None):
                calls.append(dict(params or {}))
                data = pages[min(len(calls) - 1, len(pages) - 1)]
                return FakeResp(data)

        original = m._client
        m._client = lambda: FakeClient()
        try:
            result = m.get_step_results(**kwargs)
        finally:
            m._client = original
        return json.loads(result), calls

    def test_all_mode_returns_single_array_of_all_results(self):
        result, calls = self._call([self.PAGE1, self.PAGE2], run_id=5)
        self.assertIsInstance(result, list)
        self.assertEqual(len(result), 4)
        self.assertEqual([r["id"] for r in result], [1, 2, 3, 4])
        # page walk: page=1 then page=2, run forwarded
        self.assertEqual(calls[0].get("page"), 1)
        self.assertEqual(calls[1].get("page"), 2)
        self.assertEqual(calls[0].get("run"), 5)
        self.assertEqual(calls[1].get("run"), 5)

    def test_legacy_page_returns_single_page_dict(self):
        result, calls = self._call([self.PAGE1, self.PAGE2], run_id=5, page=2)
        self.assertIsInstance(result, dict)
        self.assertEqual(len(result["results"]), 2)
        self.assertIn("next", result)
        self.assertIn("previous", result)
        self.assertEqual(calls[0].get("page"), 2)

    def test_page_size_truncates_combined_array(self):
        result, calls = self._call([self.PAGE1, self.PAGE2], run_id=5, page_size=2)
        self.assertIsInstance(result, list)
        self.assertEqual([r["id"] for r in result], [1, 2])

    def test_status_filter_forwarded(self):
        result, calls = self._call([self.PAGE1, self.PAGE2], run_id=5, status="passed")
        self.assertEqual(calls[0].get("status"), "passed")
        self.assertEqual(calls[1].get("status"), "passed")


class FindingFilterTest(TestCase):
    """FindingViewSet project_name REST filter + MCP get_findings optional filters."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        self.plan_a = TestPlan.objects.create(name='PA', project_name='alpha', created_by=self.user)
        self.plan_b = TestPlan.objects.create(name='PB', project_name='beta', created_by=self.user)
        self.run_a = TestRun.objects.create(plan=self.plan_a)
        self.run_b = TestRun.objects.create(plan=self.plan_b)
        self.f_a = Finding.objects.create(run=self.run_a, title='A1', description='d', category='critical')
        self.f_b = Finding.objects.create(run=self.run_b, title='B1', description='d', category='info')

    def test_rest_project_name_filter(self):
        r = self.client.get('/api/findings/', {'project_name': 'alpha'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual([f['id'] for f in r.json()['results']], [self.f_a.id])

    def test_rest_category_filter(self):
        r = self.client.get('/api/findings/', {'category': 'critical'})
        self.assertEqual([f['id'] for f in r.json()['results']], [self.f_a.id])

    def test_rest_combined_filters(self):
        r = self.client.get('/api/findings/', {'project_name': 'alpha', 'category': 'critical'})
        self.assertEqual([f['id'] for f in r.json()['results']], [self.f_a.id])
        r = self.client.get('/api/findings/', {'project_name': 'alpha', 'category': 'info'})
        self.assertEqual(r.json()['results'], [])

    def _mcp_call(self, **kwargs):
        import mcp_server.server as m
        calls = []

        class FakeResp:
            def __init__(self, data):
                self._data = data

            def raise_for_status(self):
                pass

            def json(self):
                return self._data

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get(self, url, params=None):
                calls.append(dict(params or {}))
                return FakeResp({"results": [{"id": 1}], "next": None})

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = m.get_findings(**kwargs)
        finally:
            m._client = original
        return json.loads(out), calls

    def test_mcp_forwards_project_name_only(self):
        out, calls = self._mcp_call(project_name='alpha')
        self.assertEqual(calls[0], {'project_name': 'alpha', 'page': 1})
        self.assertEqual(out, [{"id": 1}])

    def test_mcp_no_filters_returns_array(self):
        out, calls = self._mcp_call()
        self.assertEqual(calls[0], {'page': 1})
        self.assertIsInstance(out, list)

    def test_mcp_forwards_run_and_category(self):
        out, calls = self._mcp_call(run_id=7, category='critical')
        self.assertEqual(calls[0], {'run': 7, 'category': 'critical', 'page': 1})


class McpUpdateResultTest(TestCase):
    """update_step_result MCP tool: lookup + PATCH with friendly errors."""

    def _call(self, results, **kwargs):
        import mcp_server.server as m
        calls = []

        class FakeResp:
            def __init__(self, data):
                self._data = data

            def raise_for_status(self):
                pass

            def json(self):
                return self._data

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get(self, url, params=None):
                calls.append(("get", url, dict(params or {})))
                return FakeResp({"results": results, "next": None})

            def patch(self, url, json=None):
                calls.append(("patch", url, json))
                return FakeResp({"id": 7, "status": "failed", "log_message": ""})

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = m.update_step_result(**kwargs)
        finally:
            m._client = original
        return json.loads(out), calls

    def test_update_status_only(self):
        out, calls = self._call([{"id": 7}], run_id=1, step_id=2, status="failed")
        self.assertEqual(calls[0], ("get", "/api/step-results/", {"run": 1, "step": 2}))
        self.assertEqual(calls[1], ("patch", "/api/step-results/7/", {"status": "failed"}))
        self.assertEqual(out["id"], 7)

    def test_update_log_message_only(self):
        out, calls = self._call([{"id": 7}], run_id=1, step_id=2, log_message="fixed")
        self.assertEqual(calls[1], ("patch", "/api/step-results/7/", {"log_message": "fixed"}))

    def test_missing_result_returns_error(self):
        out, calls = self._call([], run_id=1, step_id=2, status="failed")
        self.assertIn("error", out)
        self.assertIn("log it first", out["error"])
        self.assertEqual(len(calls), 1)  # only the lookup GET

    def test_invalid_status_no_http(self):
        out, calls = self._call([{"id": 7}], run_id=1, step_id=2, status="bogus")
        self.assertIn("error", out)
        self.assertEqual(calls, [])


class McpDeleteResultTest(TestCase):
    """delete_step_result MCP tool: lookup + DELETE with friendly error."""

    def _call(self, results, **kwargs):
        import mcp_server.server as m
        calls = []

        class FakeResp:
            def __init__(self, data):
                self._data = data

            def raise_for_status(self):
                pass

            def json(self):
                return self._data

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get(self, url, params=None):
                calls.append(("get", url, dict(params or {})))
                return FakeResp({"results": results, "next": None})

            def delete(self, url, **kw):
                calls.append(("delete", url, kw))
                return FakeResp({})

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = m.delete_step_result(**kwargs)
        finally:
            m._client = original
        return json.loads(out), calls

    def test_delete_existing_result(self):
        out, calls = self._call([{"id": 7}], run_id=1, step_id=2)
        self.assertEqual(calls[0], ("get", "/api/step-results/", {"run": 1, "step": 2}))
        self.assertEqual(calls[1][0], "delete")
        self.assertEqual(calls[1][1], "/api/step-results/7/")
        self.assertEqual(out, {"deleted": 7})

    def test_missing_result_returns_error(self):
        out, calls = self._call([], run_id=1, step_id=2)
        self.assertIn("error", out)
        self.assertIn("No result for run 1 step 2", out["error"])
        self.assertEqual(len(calls), 1)  # only the lookup GET


class ReopenRunTest(TestCase):
    """POST /api/test-runs/{id}/reopen/ reopens a finished run for execution."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='P', created_by=self.user)
        self.step = TestStep.objects.create(
            plan=self.plan, name='S', action_description='a',
            expected_outcome='b', order_index=0,
        )
        self.run = TestRun.objects.create(
            plan=self.plan, status='completed', completed_at=timezone.now(),
        )

    def _reopen(self):
        return self.client.post(f'/api/test-runs/{self.run.id}/reopen/')

    def test_reopen_completed_run(self):
        r = self._reopen()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['status'], 'running')
        self.assertIsNone(r.json()['completed_at'])
        self.run.refresh_from_db()
        self.assertEqual(self.run.status, 'running')
        self.assertIsNone(self.run.completed_at)

    def test_reopen_is_idempotent(self):
        self.assertEqual(self._reopen().status_code, 200)
        r2 = self._reopen()
        self.assertEqual(r2.status_code, 200)
        self.assertEqual(r2.json()['status'], 'running')
        self.assertIsNone(r2.json()['completed_at'])

    def test_reopen_allows_new_results(self):
        self._reopen()
        r = self.client.post(
            '/api/step-results/',
            {'run': self.run.id, 'step': self.step.id, 'status': 'passed'},
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 201)

    def test_reopen_failed_run(self):
        self.run.status = 'failed'
        self.run.save()
        r = self._reopen()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['status'], 'running')
        self.assertIsNone(r.json()['completed_at'])


class McpReopenTest(TestCase):
    """reopen_test_run MCP tool wraps POST /api/test-runs/{id}/reopen/."""

    def _call(self, resp_data, status_code=200):
        import httpx
        import mcp_server.server as m
        calls = []

        class FakeResp:
            def __init__(self):
                self.status_code = status_code

            def raise_for_status(self):
                if self.status_code >= 400:
                    raise httpx.HTTPStatusError(
                        "error",
                        request=httpx.Request("POST", "http://t"),
                        response=httpx.Response(self.status_code),
                    )

            def json(self):
                return resp_data

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def post(self, url, json=None):
                calls.append((url, json))
                return FakeResp()

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = m.reopen_test_run(5)
        finally:
            m._client = original
        return json.loads(out), calls

    def test_reopen_success(self):
        out, calls = self._call({"id": 5, "status": "running", "completed_at": None})
        self.assertEqual(calls[0], ("/api/test-runs/5/reopen/", {}))
        self.assertEqual(out["status"], "running")

    def test_reopen_404_returns_error(self):
        out, calls = self._call({}, status_code=404)
        self.assertEqual(out, {"error": "Run 5 not found"})


class BlockedStatusTest(TestCase):
    """RunStepResult supports the 'blocked' status (model level)."""

    def setUp(self):
        self.plan = TestPlan.objects.create(name='P')
        self.step = TestStep.objects.create(
            plan=self.plan, name='S', action_description='a',
            expected_outcome='b', order_index=0,
        )
        self.run = TestRun.objects.create(plan=self.plan)

    def test_blocked_in_choices(self):
        self.assertIn(('blocked', 'Blocked'), RunStepResult.STATUS_CHOICES)

    def test_blocked_result_saves_and_displays(self):
        r = RunStepResult.objects.create(run=self.run, step=self.step, status='blocked')
        r.refresh_from_db()
        self.assertEqual(r.status, 'blocked')
        self.assertEqual(r.get_status_display(), 'Blocked')

    # ── Task 9: blocked accepted by log endpoints + MCP ──

    def test_bulk_log_accepts_blocked(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        r = self.client.post(
            '/api/step-results/bulk-log/',
            {
                'run': self.run.id,
                'results': [{'step': self.step.id, 'status': 'blocked',
                             'log_message': 'missing credential'}],
            },
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.json()['created'], 1)
        self.assertEqual(r.json()['results'][0]['status'], 'blocked')

    def test_bulk_log_invalid_status_still_400(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        r = self.client.post(
            '/api/step-results/bulk-log/',
            {'run': self.run.id, 'results': [{'step': self.step.id, 'status': 'bogus'}]},
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 400)
        for st in ('passed', 'failed', 'skipped', 'blocked'):
            self.assertIn(st, str(r.json()))

    def test_mcp_log_accepts_blocked(self):
        import mcp_server.server as m
        calls = []

        class FakeResp:
            def raise_for_status(self):
                pass

            def json(self):
                return {'id': 1, 'status': 'blocked'}

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def post(self, url, json=None):
                calls.append({'url': url, 'json': json})
                return FakeResp()

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = m.log_step_result(run_id=self.run.id, step_id=self.step.id, status='blocked')
        finally:
            m._client = original
        self.assertNotIn('error', json.loads(out))
        self.assertEqual(calls[0]['json']['status'], 'blocked')

    def test_mcp_docstrings_list_blocked(self):
        import mcp_server.server as m
        self.assertIn('blocked', m.log_step_result.__doc__)
        self.assertIn('blocked', m.bulk_log_step_results.__doc__)


class BlockedProgressTest(TestCase):
    """blocked counts as EXECUTED (not pending) in model + progress API."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='P', created_by=self.user)
        self.steps = [
            TestStep.objects.create(
                plan=self.plan, name=f'S{i}', action_description='a',
                expected_outcome='b', order_index=i,
            )
            for i in range(4)
        ]
        self.run = TestRun.objects.create(plan=self.plan)
        RunStepResult.objects.create(run=self.run, step=self.steps[0], status='blocked')
        RunStepResult.objects.create(run=self.run, step=self.steps[1], status='passed')
        RunStepResult.objects.create(run=self.run, step=self.steps[2], status='skipped')

    def test_model_properties_exclude_blocked_from_pending(self):
        self.assertEqual(self.run.blocked_steps, 1)
        self.assertEqual(self.run.pending_steps, 1)

    def test_progress_api_returns_blocked_and_corrected_pending(self):
        r = self.client.get(f'/api/test-runs/{self.run.id}/progress/')
        self.assertEqual(r.status_code, 200)
        d = r.json()
        self.assertEqual(d['blocked'], 1)
        self.assertEqual(d['pending'], 1)
        self.assertEqual(d['executed'], 3)


class FindingStepsTest(TestCase):
    """Finding.step_ids M2M: REST create + MCP create_finding forwarding."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='P', created_by=self.user)
        self.run = TestRun.objects.create(plan=self.plan)
        self.s1 = TestStep.objects.create(
            plan=self.plan, name='Step One', action_description='a',
            expected_outcome='b', order_index=0,
        )
        self.s2 = TestStep.objects.create(
            plan=self.plan, name='Step Two', action_description='a',
            expected_outcome='b', order_index=1,
        )

    def test_rest_create_with_step_ids(self):
        r = self.client.post(
            '/api/findings/',
            {
                'run': self.run.id, 'title': 'T', 'description': 'D',
                'category': 'info', 'step_ids': [self.s1.id, self.s2.id],
            },
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 201)
        self.assertEqual(sorted(r.json()['step_ids']), sorted([self.s1.id, self.s2.id]))
        self.assertEqual(
            r.json()['step_names'],
            [f'#{self.s1.id} Step One', f'#{self.s2.id} Step Two'],
        )

    def test_rest_create_unknown_step_400(self):
        r = self.client.post(
            '/api/findings/',
            {'run': self.run.id, 'title': 'T', 'description': 'D',
             'step_ids': [999999]},
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 400)

    def test_mcp_create_finding_forwards_step_ids(self):
        import mcp_server.server as m
        captured = {}

        class FakeResp:
            def raise_for_status(self):
                pass

            def json(self):
                return {'id': 1}

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def post(self, url, json=None):
                captured['json'] = json
                return FakeResp()

        original = m._client
        m._client = lambda: FakeClient()
        try:
            m.create_finding(run_id=1, title='T', description='D', step_ids=[1, 2])
        finally:
            m._client = original
        self.assertEqual(captured['json']['step_ids'], [1, 2])


class ExportFindingsStepsTest(TestCase):
    """Web export Findings sheet gains a Related Steps column."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='P', created_by=self.user)
        self.run = TestRun.objects.create(plan=self.plan)
        self.s1 = TestStep.objects.create(
            plan=self.plan, name='Step One', action_description='a',
            expected_outcome='b', order_index=0,
        )
        self.s2 = TestStep.objects.create(
            plan=self.plan, name='Step Two', action_description='a',
            expected_outcome='b', order_index=1,
        )
        self.f_linked = Finding.objects.create(run=self.run, title='Linked', description='D', category='info')
        self.f_linked.step_ids.set([self.s1, self.s2])
        self.f_none = Finding.objects.create(run=self.run, title='NoSteps', description='D', category='info')

    def _findings_rows(self):
        import io
        from openpyxl import load_workbook
        r = self.client.get(f'/run/{self.run.id}/export/')
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn('Findings', wb.sheetnames)
        ws = wb['Findings']
        headers = [c.value for c in ws[1]]
        self.assertIn('Related Steps', headers)
        col = headers.index('Related Steps')
        rows = {}
        for row in ws.iter_rows(min_row=4):
            if row[0].value is not None:
                rows[row[0].value] = row[col].value
        return rows

    def test_findings_row_lists_both_steps(self):
        rows = self._findings_rows()
        self.assertIn(f'#{self.s1.id} Step One', rows[self.f_linked.id])
        self.assertIn(f'#{self.s2.id} Step Two', rows[self.f_linked.id])

    def test_findings_row_without_steps_is_dash(self):
        rows = self._findings_rows()
        self.assertEqual(rows[self.f_none.id], '-')


class ApiExportTest(TestCase):
    """GET /api/test-runs/export/ (xlsx/csv) + export_run MCP tool."""

    def setUp(self):
        self.user = User.objects.create_user(username='testuser', password='testpass')
        self.client.login(username='testuser', password='testpass')
        self.plan = TestPlan.objects.create(name='P', created_by=self.user)
        self.steps = [
            TestStep.objects.create(
                plan=self.plan, name=f'S{i}', action_description='a',
                expected_outcome='b', order_index=i,
            )
            for i in range(3)
        ]
        self.run1 = TestRun.objects.create(plan=self.plan)
        self.run2 = TestRun.objects.create(plan=self.plan)
        RunStepResult.objects.create(run=self.run1, step=self.steps[0], status='passed')
        RunStepResult.objects.create(run=self.run1, step=self.steps[1], status='skipped')
        RunStepResult.objects.create(run=self.run2, step=self.steps[0], status='failed')
        self.finding = Finding.objects.create(
            run=self.run1, title='F', description='D', category='info',
        )
        self.finding.step_ids.set([self.steps[0]])

    def _sheet_status_values(self, wb, sheet_name):
        ws = wb[sheet_name]
        values = []
        for row in ws.iter_rows(min_row=2):
            if row[1].value is not None:
                values.append(row[1].value)
        return values

    def test_xlsx_export_multi_run(self):
        import io
        from openpyxl import load_workbook
        r = self.client.get(
            '/api/test-runs/export/',
            {'runs': f'{self.run1.id},{self.run2.id}', 'format': 'xlsx',
             'exclude_skipped': 'true'},
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r['Content-Type'].startswith('application/vnd.openxmlformats'))
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn(f'Run {self.run1.id}', wb.sheetnames)
        self.assertIn(f'Run {self.run2.id}', wb.sheetnames)
        self.assertIn('Findings', wb.sheetnames)
        self.assertIn('Summary', wb.sheetnames)
        for name in (f'Run {self.run1.id}', f'Run {self.run2.id}'):
            self.assertNotIn('Skipped', self._sheet_status_values(wb, name))

    def test_csv_export(self):
        r = self.client.get(
            '/api/test-runs/export/',
            {'runs': f'{self.run1.id},{self.run2.id}', 'format': 'csv',
             'exclude_skipped': 'true'},
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r['Content-Type'].startswith('text/csv'))
        lines = r.content.decode().strip().splitlines()
        self.assertEqual(
            lines[0], 'run_id,step_id,step_name,status,log_message,created_at',
        )
        statuses = [ln.split(',')[3] for ln in lines[1:]]
        self.assertNotIn('skipped', statuses)
        self.assertEqual(len(statuses), 2)  # 1 passed + 1 failed

    def test_missing_runs_400(self):
        r = self.client.get('/api/test-runs/export/', {'format': 'xlsx'})
        self.assertEqual(r.status_code, 400)

    def test_bad_format_400(self):
        r = self.client.get(
            '/api/test-runs/export/',
            {'runs': str(self.run1.id), 'format': 'pdf'},
        )
        self.assertEqual(r.status_code, 400)

    def test_api_key_auth(self):
        admin = User.objects.create_user(
            username='admin', password='adminpass', is_staff=True, is_superuser=True,
        )
        _instance, raw_key = APIKey.objects.create_key(name='export-key')
        r = self.client.get(
            '/api/test-runs/export/',
            {'runs': str(self.run1.id), 'format': 'csv'},
            HTTP_AUTHORIZATION=f'Api-Key {raw_key}',
        )
        self.assertEqual(r.status_code, 200)

    def test_mcp_export_run(self):
        import base64
        import mcp_server.server as m

        payload = b'file-bytes'
        captured = {}

        class FakeResp:
            status_code = 200
            content = payload
            headers = {'content-type': 'text/csv'}

            def raise_for_status(self):
                pass

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get(self, url, params=None):
                captured['url'] = url
                captured['params'] = dict(params or {})
                return FakeResp()

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = json.loads(m.export_run('1,2', exclude_skipped=True))
        finally:
            m._client = original
        self.assertEqual(captured['url'], '/api/test-runs/export/')
        self.assertEqual(captured['params']['runs'], '1,2')
        self.assertEqual(base64.b64decode(out['content_base64']), payload)
        self.assertEqual(out['size_bytes'], len(payload))
        self.assertEqual(out['content_type'], 'text/csv')


class McpProjectSummaryTest(TestCase):
    """get_project_summary assembles a project report from REST data."""

    def _call(self, project_name):
        import mcp_server.server as m

        def page(items):
            return {"count": len(items), "next": None, "previous": None,
                    "results": items}

        def progress(run_id, passed, failed, skipped, blocked, total):
            return {
                "run_id": run_id, "status": "completed", "total_steps": total,
                "executed": passed + failed + skipped + blocked,
                "passed": passed, "failed": failed, "skipped": skipped,
                "blocked": blocked,
                "pending": total - (passed + failed + skipped + blocked),
                "pending_step_ids": [], "sections": [],
            }

        routes = {
            "/api/test-plans/": page([
                {"id": 27, "name": "Main plan", "project_name": "bonifacil", "total_steps": 38},
                {"id": 28, "name": "Decoy", "project_name": "bonifacilio", "total_steps": 5},
            ]),
            "/api/test-runs/": page([
                {"id": 128, "plan": 27, "status": "completed"},
                {"id": 130, "plan": 27, "status": "running"},
            ]),
            "/api/test-runs/128/progress/": progress(128, 19, 0, 19, 0, 38),
            "/api/test-runs/130/progress/": progress(130, 0, 0, 10, 0, 38),
            "/api/findings/": page([
                {"id": 1, "run": 128, "category": "critical"},
                {"id": 2, "run": 130, "category": "info"},
                {"id": 3, "run": 999, "category": "info"},
            ]),
            "/api/step-results/": page([
                {"id": 100, "run": 128},
                {"id": 101, "run": 130},
            ]),
            "/api/incidents/": page([
                {"id": 1, "run_step_result": 100, "severity": "high"},
                {"id": 2, "run_step_result": 555, "severity": "low"},
            ]),
        }
        calls = []

        class FakeResp:
            def __init__(self, data):
                self._data = data

            def raise_for_status(self):
                pass

            def json(self):
                return self._data

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get(self, url, params=None):
                calls.append((url, dict(params or {})))
                return FakeResp(routes[url])

        original = m._client
        m._client = lambda: FakeClient()
        try:
            out = json.loads(m.get_project_summary(project_name))
        finally:
            m._client = original
        return out, calls

    def test_summary_shape_and_math(self):
        out, calls = self._call("bonifacil")
        self.assertEqual(out["project_name"], "bonifacil")
        # exact, case-insensitive plan match (decoy 'bonifacilio' excluded)
        self.assertEqual(out["plans"], [{"id": 27, "name": "Main plan", "total_steps": 38}])
        runs = {r["run_id"]: r for r in out["runs"]}
        self.assertEqual(set(runs), {128, 130})
        self.assertEqual(runs[128]["passed"], 19)
        self.assertEqual(runs[128]["skipped"], 19)
        self.assertEqual(runs[128]["blocked"], 0)
        self.assertEqual(runs[128]["pending"], 0)
        self.assertEqual(runs[128]["pass_rate"], 1.0)
        self.assertIsNone(runs[130]["pass_rate"])  # 0 passed + 0 failed
        self.assertEqual(out["findings_by_category"], {"critical": 1, "info": 1})
        self.assertEqual(
            out["incidents_by_severity"],
            {"low": 0, "medium": 0, "high": 1, "critical": 0},
        )

    def test_unknown_project_returns_empty_not_error(self):
        out, calls = self._call("no-such-project")
        self.assertEqual(out, {
            "project_name": "no-such-project",
            "plans": [],
            "runs": [],
            "findings_by_category": {},
            "incidents_by_severity": {},
        })
