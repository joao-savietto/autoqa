from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth import views as auth_views, login as auth_login
from django.contrib.auth.models import User
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.utils.timezone import localtime, now as tz_now
from django.core.paginator import Paginator

import environ
from datetime import datetime

env = environ.Env()
environ.Env.read_env()

from .models import TestPlan, TestRun, RunStepResult, TestStep, Incident, Finding
from rest_framework_api_key.models import APIKey


def login_view(request):
    if not User.objects.exists():
        return redirect('register')
    return auth_views.LoginView.as_view(template_name='login.html')(request)


def register_view(request):
    if User.objects.exists():
        return redirect('login')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        password_confirm = request.POST.get('password_confirm', '')

        errors = []
        if not username:
            errors.append('Username is required.')
        if not password:
            errors.append('Password is required.')
        if password != password_confirm:
            errors.append('Passwords do not match.')
        if username and User.objects.filter(username=username).exists():
            errors.append('Username already taken.')

        if not errors:
            user = User.objects.create_user(username=username, password=password)
            user.is_superuser = True
            user.is_staff = True
            user.save()
            auth_login(request, user, backend='django.contrib.auth.backends.ModelBackend')
            messages.success(request, 'Account created successfully. Welcome to AutoQA!')
            return redirect('dashboard')

        return render(request, 'register.html', {'errors': errors})

    return render(request, 'register.html', {'errors': []})


def logout_view(request):
    return auth_views.LogoutView.as_view(next_page='login')(request)


@login_required
def dashboard(request):
    plans = TestPlan.objects.all().prefetch_related('testruns').order_by('-created_at')
    return render(request, 'dashboard.html', {'plans': plans})


@login_required
def plan_detail(request, plan_id):
    plan = get_object_or_404(TestPlan, id=plan_id)
    steps_qs = plan.teststeps.filter(active=True).order_by('order_index')
    section = request.GET.get('section', '')
    if section:
        steps_qs = steps_qs.filter(section=section)

    paginator = Paginator(steps_qs, 20)
    page_number = request.GET.get('page', 1)
    steps = paginator.get_page(page_number)

    sections = plan.teststeps.filter(active=True).exclude(section='').values_list(
        'section', flat=True
    ).distinct().order_by('section')

    runs = plan.testruns.order_by('-started_at')[:20]
    return render(request, 'plan_detail.html', {
        'plan': plan, 'steps': steps, 'runs': runs, 'sections': sections,
        'current_section': section,
    })


@login_required
def plan_detail_steps(request, plan_id):
    """HTMX endpoint: returns a page of test steps as a partial."""
    plan = get_object_or_404(TestPlan, id=plan_id)
    steps_qs = plan.teststeps.filter(active=True).order_by('order_index')
    section = request.GET.get('section', '')
    if section:
        steps_qs = steps_qs.filter(section=section)

    paginator = Paginator(steps_qs, 20)
    page_number = request.GET.get('page', 1)
    steps = paginator.get_page(page_number)

    return render(request, 'partials/steps_table.html', {
        'plan': plan, 'steps': steps,
    })


@login_required
def run_detail(request, run_id):
    run = get_object_or_404(TestRun, id=run_id)
    results = run.runstepresults.select_related('step').order_by('step__order_index')
    incidents = Incident.objects.filter(run_step_result__run=run).distinct()
    findings = run.findings.all()
    active_tab = request.GET.get('tab', '')
    valid_tabs = ('steps', 'failed', 'skipped', 'incidents', 'findings')
    if active_tab not in valid_tabs:
        active_tab = ''
    return render(request, 'run_detail.html', {
        'run': run, 'results': results,
        'incidents': incidents, 'findings': findings,
        'active_tab': active_tab,
    })


@login_required
def step_edit(request, step_id):
    step = get_object_or_404(TestStep, id=step_id)
    if request.method == 'POST':
        step.name = request.POST.get('name', step.name)
        step.action_description = request.POST.get('action_description', step.action_description)
        step.preconditions = request.POST.get('preconditions', step.preconditions)
        step.expected_outcome = request.POST.get('expected_outcome', step.expected_outcome)
        step.active = request.POST.get('active') == 'on'
        step.save()
        return redirect('plan_detail', plan_id=step.plan.id)
    return render(request, 'step_edit.html', {'step': step})


@login_required
def step_delete(request, step_id):
    step = get_object_or_404(TestStep, id=step_id)
    plan_id = step.plan.id
    step.delete()
    return redirect('plan_detail', plan_id=plan_id)


@login_required
def plan_delete(request, plan_id):
    plan = get_object_or_404(TestPlan, id=plan_id)
    plan.delete()
    return redirect('dashboard')


def chrome_connection_view(request):
    """Return the Chrome DevTools connection string."""
    chrome_host = env('CHROME_HOST', default='localhost')
    chrome_port = env('CHROME_DEBUG_PORT', default='9222')
    return JsonResponse({
        'connection_string': f'http://{chrome_host}:{chrome_port}',
        'host': chrome_host,
        'port': int(chrome_port),
    })


@login_required
def api_keys_view(request):
    """UI page to view and manage API keys for agent authentication."""
    keys = APIKey.objects.all().order_by('-created')
    return render(request, 'api_keys.html', {'keys': keys})


def _styled_header(ws, headers, font_color="FFFFFF", fill_color="1E293B"):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    header_font = Font(name="Segoe UI", size=11, bold=True, color=font_color)
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="334155"),
        right=Side(style="thin", color="334155"),
        bottom=Side(style="thin", color="334155"),
    )
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border


def _styled_row(ws, row_idx, values, is_alt=False):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    bg_color = "E8EEF6" if is_alt else "FFFFFF"
    row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    cell_font = Font(name="Segoe UI", size=10)
    cell_align = Alignment(vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = cell_font
        cell.fill = row_fill
        cell.alignment = cell_align
        cell.border = thin_border


def _status_font(status):
    from openpyxl.styles import Font
    status_lower = status.lower() if status else ""
    if status_lower == "passed":
        return Font(name="Segoe UI", size=10, bold=True, color="059669")
    elif status_lower == "failed":
        return Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    elif status_lower == "skipped":
        return Font(name="Segoe UI", size=10, color="6B7280")
    elif status_lower in ("completed", "active"):
        return Font(name="Segoe UI", size=10, bold=True, color="059669")
    elif status_lower == "running":
        return Font(name="Segoe UI", size=10, bold=True, color="D97706")
    elif status_lower == "pending":
        return Font(name="Segoe UI", size=10, color="6B7280")
    return Font(name="Segoe UI", size=10)


def _severity_font(severity):
    from openpyxl.styles import Font
    sev = severity.lower() if severity else ""
    if sev == "high":
        return Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    elif sev == "medium":
        return Font(name="Segoe UI", size=10, bold=True, color="D97706")
    return Font(name="Segoe UI", size=10, color="3B82F6")


def _category_font(category):
    from openpyxl.styles import Font
    cat = category.lower() if category else ""
    colors = {"critical": "DC2626", "recommendation": "059669", "suggestion": "D97706"}
    color = colors.get(cat, "3B82F6")
    return Font(name="Segoe UI", size=10, bold=True, color=color)


def _title_row(ws, row_idx, title, subtitle=None):
    from openpyxl.styles import Font, PatternFill, Alignment
    title_font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    title_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    title_cell = ws.cell(row=row_idx, column=1, value=title)
    title_cell.font = title_font
    title_cell.alignment = Alignment(vertical="center")
    if subtitle:
        sub_cell = ws.cell(row=row_idx + 1, column=1, value=subtitle)
        sub_cell.font = subtitle_font
    for cell in ws[row_idx]:
        cell.fill = title_fill
    if subtitle:
        for cell in ws[row_idx + 1]:
            cell.fill = title_fill


@login_required
def export_steps_xlsx(request, plan_id):
    plan = get_object_or_404(TestPlan, id=plan_id)
    steps = plan.teststeps.all().order_by('order_index')

    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Steps"
    ws.sheet_properties.tabColor = "6366F1"

    _title_row(ws, 1, plan.name, f"Project: {plan.project_name or 'N/A'}  |  Type: {plan.plan_type.upper()}  |  Generated: {localtime(tz_now()).strftime('%Y-%m-%d %H:%M')}")

    header_row = 4
    headers = ["#", "Name", "Action Description", "Preconditions", "Expected Outcome", "Status", "Created", "Updated"]
    _styled_header(ws, headers)
    ws.row_dimensions[header_row].height = 30

    for row_idx, step in enumerate(steps, header_row + 1):
        is_alt = (row_idx - header_row) % 2 == 0
        values = [
            step.order_index,
            step.name,
            step.action_description,
            step.preconditions or "—",
            step.expected_outcome,
            "Active" if step.active else "Inactive",
            localtime(step.created_at).strftime("%Y-%m-%d %H:%M"),
            localtime(step.updated_at).strftime("%Y-%m-%d %H:%M"),
        ]
        _styled_row(ws, row_idx, values, is_alt)
        status_cell = ws.cell(row=row_idx, column=6)
        status_cell.font = _status_font("Active" if step.active else "Inactive")
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [6, 28, 40, 35, 40, 10, 16, 16]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    filename = f"test_steps_{plan.name.replace(' ', '_')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def export_run_results_xlsx(request, run_id):
    run = get_object_or_404(TestRun, id=run_id)
    results = run.runstepresults.select_related('step').order_by('step__order_index')
    incidents = Incident.objects.filter(run_step_result__run=run).distinct()
    findings = run.findings.all()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    wb = Workbook()

    # Sheet 1: Step Results
    ws = wb.active
    ws.title = "Step Results"
    ws.sheet_properties.tabColor = "6366F1"

    _title_row(ws, 1, f"Run #{run.id} — {run.plan.name}",
               f"Status: {run.get_status_display()}  |  Started: {localtime(run.started_at).strftime('%Y-%m-%d %H:%M')}  |  Agent: {run.agent_id or 'N/A'}")

    if run.completed_at:
        from openpyxl.styles import Font as _Font, PatternFill as _Fill
        ws.cell(row=3, column=1, value=f"Completed: {localtime(run.completed_at).strftime('%Y-%m-%d %H:%M')}").font = _Font(name="Segoe UI", size=10, italic=True, color="64748B")
        for cell in ws[3]:
            cell.fill = _Fill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    # Summary stats row
    stats_row = 5
    from openpyxl.styles import Font, PatternFill, Border, Side
    stats_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    stats_font = Font(name="Segoe UI", size=10, bold=True, color="334155")
    stats_border = Border(bottom=Side(style="medium", color="6366F1"))
    stats_labels = ["Total", "Passed", "Failed", "Skipped", "Pending"]
    stats_values = [run.total_steps, run.passed_steps, run.failed_steps, run.skipped_steps, run.pending_steps]
    stats_colors = ["334155", "059669", "DC2626", "6B7280", "D97706"]
    for i, (label, val, color) in enumerate(zip(stats_labels, stats_values, stats_colors), 1):
        lbl_cell = ws.cell(row=stats_row, column=i, value=label)
        lbl_cell.font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
        lbl_cell.alignment = Alignment(horizontal="center", vertical="bottom")
        lbl_cell.fill = stats_fill
        lbl_cell.border = stats_border
        val_cell = ws.cell(row=stats_row + 1, column=i, value=val)
        val_cell.font = Font(name="Segoe UI", size=14, bold=True, color=color)
        val_cell.alignment = Alignment(horizontal="center", vertical="top")
        val_cell.fill = stats_fill
        val_cell.border = stats_border

    header_row = stats_row + 3
    headers = ["#", "Step Name", "Status", "Action Description", "Expected Outcome", "Log Message", "Executed At"]
    _styled_header(ws, headers)
    ws.row_dimensions[header_row].height = 30

    for row_idx, result in enumerate(results, header_row + 1):
        is_alt = (row_idx - header_row) % 2 == 0
        values = [
            result.step.order_index,
            result.step.name,
            result.get_status_display(),
            result.step.action_description,
            result.step.expected_outcome,
            result.log_message or "—",
            localtime(result.created_at).strftime("%Y-%m-%d %H:%M"),
        ]
        _styled_row(ws, row_idx, values, is_alt)
        status_cell = ws.cell(row=row_idx, column=3)
        status_cell.font = _status_font(result.status)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [6, 30, 12, 40, 35, 45, 16]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = width

    # Sheet 2: Failed Steps
    failed_results = [r for r in results if r.status == 'failed']
    if failed_results:
        ws2 = wb.create_sheet("Failed")
        ws2.sheet_properties.tabColor = "DC2626"
        _title_row(ws2, 1, f"Failed Steps — Run #{run.id}", f"Total: {len(failed_results)} failed out of {run.total_steps} steps")
        header_row = 3
        headers = ["#", "Step Name", "Status", "Action Description", "Expected Outcome", "Log Message", "Executed At"]
        _styled_header(ws2, headers)
        ws2.row_dimensions[header_row].height = 30
        for row_idx, result in enumerate(failed_results, header_row + 1):
            is_alt = (row_idx - header_row) % 2 == 0
            values = [
                result.step.order_index,
                result.step.name,
                result.get_status_display(),
                result.step.action_description,
                result.step.expected_outcome,
                result.log_message or "—",
                localtime(result.created_at).strftime("%Y-%m-%d %H:%M"),
            ]
            _styled_row(ws2, row_idx, values, is_alt)
            status_cell = ws2.cell(row=row_idx, column=3)
            status_cell.font = _status_font(result.status)
            status_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        fail_widths = [6, 30, 12, 40, 35, 45, 16]
        for col_idx, width in enumerate(fail_widths, 1):
            ws2.column_dimensions[chr(64 + col_idx)].width = width

    # Sheet 3: Skipped Steps
    skipped_results = [r for r in results if r.status == 'skipped']
    if skipped_results:
        ws3 = wb.create_sheet("Skipped")
        ws3.sheet_properties.tabColor = "6B7280"
        _title_row(ws3, 1, f"Skipped Steps — Run #{run.id}", f"Total: {len(skipped_results)} skipped out of {run.total_steps} steps")
        header_row = 3
        headers = ["#", "Step Name", "Status", "Action Description", "Expected Outcome", "Log Message", "Executed At"]
        _styled_header(ws3, headers)
        ws3.row_dimensions[header_row].height = 30
        for row_idx, result in enumerate(skipped_results, header_row + 1):
            is_alt = (row_idx - header_row) % 2 == 0
            values = [
                result.step.order_index,
                result.step.name,
                result.get_status_display(),
                result.step.action_description,
                result.step.expected_outcome,
                result.log_message or "—",
                localtime(result.created_at).strftime("%Y-%m-%d %H:%M"),
            ]
            _styled_row(ws3, row_idx, values, is_alt)
            status_cell = ws3.cell(row=row_idx, column=3)
            status_cell.font = _status_font(result.status)
            status_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws3.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        skip_widths = [6, 30, 12, 40, 35, 45, 16]
        for col_idx, width in enumerate(skip_widths, 1):
            ws3.column_dimensions[chr(64 + col_idx)].width = width

    # Sheet 4: Incidents
    if incidents:
        ws4 = wb.create_sheet("Incidents")
        ws4.sheet_properties.tabColor = "DC2626"
        _title_row(ws4, 1, f"Incidents — Run #{run.id}", f"Total: {len(incidents)} issues found")
        header_row = 3
        headers = ["#", "Summary", "Severity", "Reproduction Steps", "Resolved", "Created"]
        _styled_header(ws4, headers)
        ws4.row_dimensions[header_row].height = 30
        for row_idx, incident in enumerate(incidents, header_row + 1):
            is_alt = (row_idx - header_row) % 2 == 0
            values = [
                incident.id,
                incident.summary,
                incident.get_severity_display(),
                incident.reproduction_steps,
                "Yes" if incident.resolved else "No",
                localtime(incident.created_at).strftime("%Y-%m-%d %H:%M"),
            ]
            _styled_row(ws4, row_idx, values, is_alt)
            ws4.cell(row=row_idx, column=3).font = _severity_font(incident.severity)
            ws4.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center", vertical="center")
            ws4.cell(row=row_idx, column=5).alignment = Alignment(horizontal="center", vertical="center")
            ws4.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        inc_widths = [6, 35, 12, 50, 10, 16]
        for col_idx, width in enumerate(inc_widths, 1):
            ws4.column_dimensions[chr(64 + col_idx)].width = width

    # Sheet 5: Findings
    if findings:
        ws5 = wb.create_sheet("Findings")
        ws5.sheet_properties.tabColor = "D97706"
        _title_row(ws5, 1, f"Findings — Run #{run.id}", f"Total: {len(findings)} findings")
        header_row = 3
        headers = ["#", "Title", "Category", "Description", "Created"]
        _styled_header(ws5, headers)
        ws5.row_dimensions[header_row].height = 30
        for row_idx, finding in enumerate(findings, header_row + 1):
            is_alt = (row_idx - header_row) % 2 == 0
            values = [
                finding.id,
                finding.title,
                finding.get_category_display(),
                finding.description,
                localtime(finding.created_at).strftime("%Y-%m-%d %H:%M"),
            ]
            _styled_row(ws5, row_idx, values, is_alt)
            ws5.cell(row=row_idx, column=3).font = _category_font(finding.category)
            ws5.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center", vertical="center")
            ws5.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        find_widths = [6, 30, 16, 55, 16]
        for col_idx, width in enumerate(find_widths, 1):
            ws5.column_dimensions[chr(64 + col_idx)].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    filename = f"run_{run.id}_{run.plan.name.replace(' ', '_')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
